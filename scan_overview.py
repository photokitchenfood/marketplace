#!/usr/bin/env python3
"""Lightweight read-only scan: for every slide, print slide layout name,
title-placeholder text, any small AUTO_SHAPE badge texts (layout code /
category label / banner), and whether a picture placeholder is present.
Used to map out the shot-list structure across the whole deck.
"""
import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

path = sys.argv[1]
prs = Presentation(path)

for i, slide in enumerate(prs.slides):
    title_text = None
    badges = []
    has_pic_placeholder = False
    other_placeholder_texts = []

    for shape in slide.shapes:
        if getattr(shape, "is_placeholder", False):
            try:
                ptype = shape.placeholder_format.type
            except Exception:
                ptype = None
            if str(ptype).startswith("TITLE"):
                title_text = shape.text_frame.text if shape.has_text_frame else None
            elif str(ptype).startswith("PICTURE"):
                has_pic_placeholder = True
            else:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    t = shape.text_frame.text.strip().replace("\n", " | ")[:40]
                    other_placeholder_texts.append(t)
        elif shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            if shape.has_text_frame and shape.text_frame.text.strip():
                w = Emu(shape.width).inches if shape.width else 0
                h = Emu(shape.height).inches if shape.height else 0
                # small badge-like shapes only
                if w < 4.5 and h < 0.6:
                    badges.append(shape.text_frame.text.strip().replace("\n", " "))

    print(f"[{i+1:>3}] layout={slide.slide_layout.name:<35} title={title_text!r:<30} pic={has_pic_placeholder} badges={badges} other={other_placeholder_texts}")
